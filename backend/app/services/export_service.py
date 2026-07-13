"""
Export service – generates Excel files for customer data downloads.
"""
from __future__ import annotations

import io
import logging
from datetime import datetime
from typing import Any, Dict, List, Optional

from sqlalchemy import text
from sqlalchemy.orm import Session

from app.services.channel_classification import add_channel_filter
from app.services.region_filter import REGION_OPTIONS

logger = logging.getLogger(__name__)

_EXPORT_COLUMNS = [
    ("customer_name",         "客户名称",     30),
    ("industry",              "行业",         15),
    ("region",                "区域",         12),
    ("owner_name",            "负责人",       15),
    ("campaign_tag",          "专项",         15),
    ("purchase_stage",        "采购阶段",     15),
    ("forecast_type",         "预测类别",     12),
    ("role_coverage",         "关键角色覆盖", 14),
    ("intent_level",          "合作意向",     10),
    ("intent_score",          "意向分",       10),
    ("interaction_count_30d", "近30天互动",   12),
    ("interaction_count_total","总互动次数",  12),
    ("last_interaction_time", "最近互动时间", 20),
    ("last_interaction_channel","最近互动渠道",14),
    ("active_opp_count",      "在途商机数",   12),
    ("active_opp_amount",     "在途金额(万)", 14),
    ("contact_count",         "联系人数",     10),
    ("won_amount",            "已成交金额(万)",14),
]


def export_customers_excel(
    db: Session,
    *,
    keyword: Optional[str] = None,
    industry: Optional[str] = None,
    region: Optional[str] = None,
    region_keyword: Optional[str] = None,
    owner: Optional[str] = None,
    owner_keyword: Optional[str] = None,
    stage: Optional[str] = None,
    intent_level: Optional[str] = None,
    interaction_min: Optional[int] = None,
    interaction_period: int = 30,
    attribute: Optional[str] = None,
    channel: Optional[str] = None,
    sort_by: str = "intent_score",
    sort_order: str = "DESC",
) -> bytes:
    """Query dws_customer_360 and return .xlsx bytes."""
    where_parts = ["1=1"]
    params: Dict[str, Any] = {}

    if keyword:
        where_parts.append("customer_name LIKE :keyword")
        params["keyword"] = f"%{keyword}%"
    if industry:
        where_parts.append("industry = :industry")
        params["industry"] = industry
    if region:
        if region == "其他":
            placeholders = []
            for index, value in enumerate(item for item in REGION_OPTIONS if item != "其他"):
                key = f"standard_region_{index}"
                placeholders.append(f":{key}")
                params[key] = value
            where_parts.append(
                "(region IS NULL OR region = '' "
                f"OR region NOT IN ({', '.join(placeholders)}))"
            )
        else:
            where_parts.append("region = :region")
            params["region"] = region
    elif region_keyword:
        if region_keyword.strip() == "其他":
            placeholders = []
            for index, value in enumerate(item for item in REGION_OPTIONS if item != "其他"):
                key = f"standard_region_{index}"
                placeholders.append(f":{key}")
                params[key] = value
            where_parts.append(
                "(region IS NULL OR region = '' "
                f"OR region NOT IN ({', '.join(placeholders)}))"
            )
        else:
            where_parts.append("region LIKE :region_keyword")
            params["region_keyword"] = f"%{region_keyword}%"
    if owner:
        where_parts.append("owner_name = :owner")
        params["owner"] = owner
    elif owner_keyword:
        where_parts.append("owner_name LIKE :owner_keyword")
        params["owner_keyword"] = f"%{owner_keyword}%"
    if stage:
        where_parts.append("purchase_stage = :stage")
        params["stage"] = stage
    if intent_level:
        where_parts.append("intent_level = :intent_level")
        params["intent_level"] = intent_level
    if interaction_min is not None and interaction_min > 0:
        # 使用子查询动态计算指定时间范围内的互动次数
        where_parts.append("""
            customer_name IN (
                SELECT customer_name
                FROM dws_interaction_detail
                WHERE event_time >= DATE_SUB(NOW(), INTERVAL :period DAY)
                GROUP BY customer_name
                HAVING COUNT(*) >= :interaction_min
            )
        """)
        params["interaction_min"] = interaction_min
        params["period"] = interaction_period
    if attribute == "heavy":
        where_parts.append("attribute = :heavy_attribute")
        params["heavy_attribute"] = "H"
    elif attribute == "non_heavy":
        where_parts.append("(attribute IS NULL OR attribute != :heavy_attribute)")
        params["heavy_attribute"] = "H"
    add_channel_filter(
        where_parts,
        params,
        column="last_interaction_channel",
        channel=channel,
    )

    where_sql = " AND ".join(where_parts)
    allowed = {"customer_name","industry","intent_score","interaction_count_30d",
               "interaction_count_total","last_interaction_time","active_opp_amount","won_amount"}
    if sort_by not in allowed:
        sort_by = "intent_score"
    order_dir = "ASC" if sort_order.upper() == "ASC" else "DESC"

    sql = text(
        f"SELECT * FROM dws_customer_360 WHERE {where_sql} "
        f"ORDER BY {sort_by} {order_dir} LIMIT 20000"
    )
    rows = db.execute(sql, params).mappings().all()

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "客户360"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")

    # Header row
    for col_idx, (_, header, width) in enumerate(_EXPORT_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    # Data rows
    for row_idx, item in enumerate(rows, 2):
        for col_idx, (key, _, _) in enumerate(_EXPORT_COLUMNS, 1):
            value = item.get(key, "")
            if isinstance(value, datetime):
                value = value.strftime("%Y-%m-%d %H:%M")
            elif value is None:
                value = ""
            ws.cell(row=row_idx, column=col_idx, value=value)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
