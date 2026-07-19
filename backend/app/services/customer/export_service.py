"""
Export service – generates Excel files for customer data downloads.
"""
from __future__ import annotations

import io
import logging
from datetime import datetime
from typing import List, Optional

from sqlalchemy.orm import Session

from app.services.customer.customer_service import get_customer_list
from app.services.customer.key_account_query import fetch_key_accounts

logger = logging.getLogger(__name__)

_EXPORT_COLUMNS = [
    ("customer_name",         "客户名称",     30),
    ("industry",              "行业",         15),
    ("region",                "区域",         12),
    ("owner_name",            "负责人",       15),
    ("campaign_tag",          "专项",         15),
    ("purchase_stage",        "采购阶段",     15),
    ("forecast_type",         "预测类别",     12),
    ("role_coverage",         "关键角色覆盖", 14),
    ("intent_level",          "合作意向",     10),
    ("intent_score",          "意向分",       10),
    ("interaction_count_30d", "近30天互动",   12),
    ("interaction_count_total","总互动次数",  12),
    ("last_interaction_time", "最近互动时间", 20),
    ("last_interaction_channel","最近互动渠道",14),
    ("active_opp_count",      "在途商机数",   12),
    ("active_opp_amount",     "在途金额(万)", 14),
    ("contact_count",         "联系人数",     10),
    ("won_amount",            "已成交金额(万)",14),
]


def export_customers_excel(
    db: Session,
    *,
    keyword: Optional[List[str]] = None,
    special_project: Optional[List[str]] = None,
    industry: Optional[List[str]] = None,
    region: Optional[List[str]] = None,
    region_keyword: Optional[str] = None,
    owner: Optional[List[str]] = None,
    owner_keyword: Optional[str] = None,
    stage: Optional[str] = None,
    intent_level: Optional[str] = None,
    interaction_min: Optional[int] = None,
    interaction_period: int = 30,
    attribute: Optional[str] = None,
    channel: Optional[List[str]] = None,
    sort: Optional[str] = None,
) -> bytes:
    """Query the selected customer source and return .xlsx bytes."""
    special_project = special_project or []
    if special_project == ["重客"]:
        rows = fetch_key_accounts(
            db,
            keyword=keyword,
            industry=industry,
            region=region,
            region_keyword=region_keyword,
            owner=owner,
            owner_keyword=owner_keyword,
            stage=stage,
            intent_level=intent_level,
            interaction_min=interaction_min,
            interaction_period=interaction_period,
            channel=channel,
            sort=sort,
            limit=20000,
        )
    else:
        items = get_customer_list(
            db,
            keyword=keyword,
            special_project=special_project,
            industry=industry,
            region=region,
            region_keyword=region_keyword,
            owner=owner,
            owner_keyword=owner_keyword,
            stage=stage,
            intent_level=intent_level,
            interaction_min=interaction_min,
            interaction_period=interaction_period,
            attribute=attribute,
            channel=channel,
            sort=sort,
            page=1,
            page_size=20000,
        )["items"]
        rows = items

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "客户360"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")

    # Header row
    for col_idx, (_, header, width) in enumerate(_EXPORT_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    # Data rows
    for row_idx, item in enumerate(rows, 2):
        for col_idx, (key, _, _) in enumerate(_EXPORT_COLUMNS, 1):
            value = item.get(key, "")
            if isinstance(value, datetime):
                value = value.strftime("%Y-%m-%d %H:%M")
            elif value is None:
                value = ""
            ws.cell(row=row_idx, column=col_idx, value=value)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
